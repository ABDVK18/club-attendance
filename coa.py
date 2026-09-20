from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Initialize presentation with 16:9 widescreen
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Dark modern tech palette
BG_COLOR = RGBColor(15, 23, 42)       # Slate 900
TEXT_MAIN = RGBColor(248, 250, 252)   # Slate 50
TEXT_MUTED = RGBColor(148, 163, 184)  # Slate 400
ACCENT_BLUE = RGBColor(56, 189, 248)  # Sky 400

slides_data = [
    {
        "title": "Edge-AI Interactive Robotic Node",
        "subtitle": "Low-Latency Thin-Client Architecture with Hardware-Accelerated I2S & SPI Pipelines\n\nDomain: AI in Hardware (Edge AI)\nEvent: HACKFEST-2026 | Department of ECE, SRMIST\nTeam: [Team Name] | Members: Member 1, Member 2, Member 3 (Dept of ECE)",
        "bullets": []
    },
    {
        "title": "Abstract",
        "subtitle": "",
        "bullets": [
            "System Overview: A modular, low-power edge robotics architecture bridging physical interaction with edge/cloud intelligence.",
            "Distributed Paradigm: Eliminates thermal and power penalties of local edge inference by offloading models and dedicating local silicon to deterministic I/O.",
            "Hardware Integration: Custom carrier PCB integrating an all-digital I2S acoustic interface (INMP441/MAX98357A) and DMA-driven 40 MHz SPI display (ST7789).",
            "Primary Value: Demonstrates over 80% BOM reduction compared to monolithic compute systems while maintaining sub-second conversational latency."
        ]
    },
    {
        "title": "Problem Statement",
        "subtitle": "",
        "bullets": [
            "Compute & Thermal Limits: Local NLP and computer vision inference on edge SoCs cause severe thermal throttling, high idle current, and battery drain.",
            "Signal Integrity Issues: Conventional analog electret microphones and PWM audio output suffer from heavy EMI and ground bounce from actuator switching.",
            "Bus Bandwidth Bottlenecks: Standard I2C display buses (~400 kHz) achieve <1 FPS at 240x240 RGB565, causing CPU stalls and severe visual stutter."
        ]
    },
    {
        "title": "Motivation & Need",
        "subtitle": "",
        "bullets": [
            "Democratizing Physical AI: Existing companion robots cost ₹40,000+; there is a severe market need for an expressive architecture under ₹2,000.",
            "Hardware-Software Co-Design: Utilizing deterministic bus protocols (I2S and SPI with DMA) allows low-cost processors to run high-throughput tasks without OS jitter.",
            "Target Applications: Assistive desktop nodes, laboratory monitoring telemetry, and interactive smart kiosk interfaces."
        ]
    },
    {
        "title": "Proposed Solution & Architecture",
        "subtitle": "",
        "bullets": [
            "Distributed Processing: Thin-client edge node handles real-time audio/display routing, offloading heavy reasoning to a backend server.",
            "Acoustic Front-End (I2S): INMP441 digital MEMS microphone and MAX98357A Class-D amplifier eliminate DAC jitter and analog line noise.",
            "High-Speed Display Engine: 240x240 IPS panel driven over 40 MHz SPI via Direct Memory Access (DMA), sustaining 30+ FPS animation at 0% CPU load.",
            "Custom Carrier PCB: Altium-designed PCB featuring star grounding, matched trace lengths, and isolated power rails to prevent brownouts."
        ]
    },
    {
        "title": "Outcome & Impact",
        "subtitle": "",
        "bullets": [
            "Quantitative Latency: Achieved sub-850 ms round-trip voice-to-expression pipeline over local network infrastructure.",
            "Acoustic Signal Purity: High SNR maintained by keeping digital I2S traces under 30 mm and physically isolating ground planes from inductive loads.",
            "Economic Viability: Prototype BOM constrained under ₹2,500, with a scalable roadmap to a dedicated microcontroller board at ~₹500 in volume.",
            "Practical Applicability: Immediately deployable across lab environments, education robotics kits, and smart service desks."
        ]
    },
    {
        "title": "Conclusion & Roadmap",
        "subtitle": "",
        "bullets": [
            "Project Summary: Validated that a protocol-optimized thin-client architecture delivers expressive physical interaction without expensive processors.",
            "Milestones Completed: Schematics verified, bus timing benchmarked (I2S / SPI-DMA), and power distribution simulated.",
            "Round 2 Deployment: Physical PCB population, chassis mechanical integration, and live hardware demonstration at SRM campus on Oct 8-9.",
            "Team Contact: [teamlead_email@srmist.edu.in] | GitHub: [github.com/your-repo]"
        ]
    }
]

for idx, data in enumerate(slides_data):
    slide = prs.slides.add_slide(blank_layout)
    
    # Background fill
    bg = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.fill.background()

    # Title Box
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.2))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = data["title"]
    p.font.size = Pt(32 if idx > 0 else 38)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    # Content Box
    content_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(4.5))
    ctf = content_box.text_frame
    ctf.word_wrap = True

    if data["subtitle"]:
        p_sub = ctf.paragraphs[0]
        p_sub.text = data["subtitle"]
        p_sub.font.size = Pt(20)
        p_sub.font.color.rgb = TEXT_MAIN

    for bullet_idx, bullet in enumerate(data["bullets"]):
        p_b = ctf.add_paragraph() if (bullet_idx > 0 or data["subtitle"]) else ctf.paragraphs[0]
        p_b.text = f"•  {bullet}"
        p_b.font.size = Pt(18)
        p_b.font.color.rgb = TEXT_MAIN
        p_b.space_after = Pt(14)

# File naming convention according to hackfest rules
prs.save("TeamName_AIinHardware_SRMIST.pptx")
print("Presentation generated successfully: TeamName_AIinHardware_SRMIST.pptx")